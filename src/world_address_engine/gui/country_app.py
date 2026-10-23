"""Offline country analysis, separate from the strict postal workflow."""

from __future__ import annotations

import csv
import json
import threading
from functools import partial
from pathlib import Path

from PySide6.QtCore import Qt, Signal, Slot, QUrl
from PySide6.QtGui import QDesktopServices
from PySide6.QtWidgets import (
    QAbstractItemView, QComboBox, QFileDialog, QFormLayout, QGroupBox,
    QHBoxLayout, QHeaderView, QLabel, QLineEdit, QListWidget, QMainWindow,
    QProgressBar, QPushButton, QTableWidget, QTableWidgetItem, QTextEdit,
    QVBoxLayout, QWidget,
)

from world_address_engine.gui.app import WorkerThread, _LIVE_THREADS, _thread_finished


def _discover_geo_pack() -> str:
    """Prefer the newest installed k2 Extended pack, then Compact."""
    root = Path.cwd() / 'packs'
    for profile in ('extended', 'compact'):
        for folder in sorted(root.glob(f'geonames-{profile}-global-*-k2'), reverse=True):
            db = folder / 'pack.db'
            if db.is_file():
                return str(db)
    return ''


def _input_header(path: Path, encoding: str, sheet: str | None) -> tuple[list[str], str | None]:
    from world_address_engine.io.delimited import _header_names

    if path.suffix.lower() == '.txt':
        return ['Text_1'], None
    if path.suffix.lower() == '.xlsx':
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = sheet or book.sheetnames[0]
            header = next(book[sheet].iter_rows(values_only=True), ())
        finally:
            book.close()
        return _header_names([str(value) if value is not None else '' for value in header]), sheet
    if path.suffix.lower() != '.csv':
        raise ValueError('Choose a CSV, TXT, or XLSX file.')
    with path.open(encoding=encoding, errors='surrogateescape', newline='') as stream:
        header = next(csv.reader(stream, strict=True), ())
    return _header_names(header), None


def _preview_input(path: Path, options: dict, cancel_requested=None) -> dict:
    from world_address_engine.cli import preview

    header, sheet = _input_header(path, options['encoding'], options['sheet_name'])
    selected = tuple(options['address_columns'])
    # Discover custom headers before asking the operator to map them. The CLI
    # requires a mapping when it cannot detect NAME_ADDRESS_LINE columns.
    detected = tuple(name for name in header
                     if name.startswith('NAME_ADDRESS_LINE') or name.startswith('Text_'))
    columns = selected or detected
    result = None
    if columns:
        result = preview(path, encoding=options['encoding'], address_columns=columns,
                         sheet_name=sheet, cancel_requested=cancel_requested)
    return {'header': header, 'selected': columns, 'sheet': sheet, 'preview': result}


def _read_country_export(output_dir: Path, report: dict) -> dict:
    """Read only the small published manifest/summary, always on the worker."""
    manifest = output_dir / 'Run_Manifest.json'
    merged = dict(report)
    if manifest.is_file():
        published = json.loads(manifest.read_text(encoding='utf-8-sig'))
        if published.get('job_kind') != 'country-analysis':
            raise ValueError('The output folder does not contain a country-analysis run.')
        merged.update(published)
    rows = []
    summary = output_dir / 'Country_Summary.csv'
    if summary.is_file():
        with summary.open(encoding='utf-8-sig', newline='') as stream:
            for row in csv.DictReader(stream):
                code = row['Country_Code']
                name = (row.get('Country_Name_EN') or row.get('Country_Name')
                        or row.get('Country_Name_TC') or '')
                rows.append((code, name if code else 'Unresolved', int(row['Count'])))
    else:
        rows = [(code, '', count) for code, count in
                sorted(merged.get('destination_counts', {}).items())]
    return {'report': merged, 'rows': rows}


class CountryAnalysisWindow(QMainWindow):
    """Import, preview, analyze, and resume country counts without postal approval."""

    jobFinished = Signal()
    jobFailed = Signal(str)

    def __init__(self):
        super().__init__()
        self.setWindowTitle('World Address Sorting Engine - Country Code Analysis')
        self.resize(1040, 800)
        self._thread = None
        self._worker_pending = False
        self._cancel = threading.Event()
        self._closing = False
        self._on_success = None
        self._on_error = None
        self._job_state = 'idle'
        self._last_error = ''
        self._last_report = None
        self._output_dir = None
        self._build()

    def _build(self) -> None:
        central = QWidget()
        layout = QVBoxLayout(central)
        self.inputs = QGroupBox('Address input and offline reference data')
        form = QFormLayout(self.inputs)
        self.input_path = QLineEdit()
        self.output_dir = QLineEdit()
        self.output_dir.setPlaceholderText('Defaults to <input file>-country-analysis')
        self.geo_db = QLineEdit(_discover_geo_pack())
        self.geo_db.setPlaceholderText('Choose an installed GeoNames pack; no download is needed')
        for label, edit, button_text, callback in (
            ('Address file', self.input_path, 'Choose file...', self._choose_input),
            ('Output folder', self.output_dir, 'Choose folder...', self._choose_output),
            ('GeoNames pack', self.geo_db, 'Choose pack...', self._choose_geo_db),
        ):
            row = QHBoxLayout()
            row.addWidget(edit, 1)
            button = QPushButton(button_text)
            button.clicked.connect(callback)
            row.addWidget(button)
            form.addRow(label, row)
        self.encoding = QComboBox()
        self.encoding.setEditable(True)
        self.encoding.addItems(['utf-8-sig', 'utf-8', 'big5', 'gb18030', 'cp1252'])
        form.addRow('Encoding', self.encoding)
        self.sheet_name = QLineEdit()
        self.sheet_name.setPlaceholderText('XLSX only; blank uses the first worksheet')
        form.addRow('Worksheet', self.sheet_name)
        self.address_columns = QListWidget()
        self.address_columns.setSelectionMode(QAbstractItemView.SelectionMode.MultiSelection)
        self.address_columns.setMinimumHeight(110)
        self.address_columns.setMaximumHeight(160)
        self.input_path.textChanged.connect(self.address_columns.clear)
        form.addRow('Address columns', self.address_columns)
        self.column_selection = QLabel()
        self.column_selection.setWordWrap(True)
        self.address_columns.itemSelectionChanged.connect(self._update_column_selection)
        self._update_column_selection()
        form.addRow('', self.column_selection)
        form.addRow(QLabel('Country catalogue: packaged catalogue (automatic)'))
        layout.addWidget(self.inputs)

        buttons = QHBoxLayout()
        for name, text, callback in (
            ('preview_button', 'Preview / load columns', self.preview_input),
            ('run_button', 'Analyze countries', self.start_job),
            ('resume_button', 'Resume', self.resume_job),
            ('cancel_button', 'Cancel', self.cancel_job),
        ):
            button = QPushButton(text)
            button.clicked.connect(callback)
            setattr(self, name, button)
            buttons.addWidget(button)
        self.cancel_button.setEnabled(False)
        self.progress = QProgressBar()
        self.progress.setRange(0, 1)
        buttons.addWidget(self.progress, 1)
        layout.addLayout(buttons)
        self.preview_summary = QLabel('Preview reads the input offline without displaying addresses.')
        self.preview_summary.setWordWrap(True)
        layout.addWidget(self.preview_summary)
        self.summary = QLabel('No analysis loaded. Determined, inferred, and unresolved records will appear here.')
        self.summary.setWordWrap(True)
        layout.addWidget(self.summary)
        self.counts = QTableWidget(0, 3)
        self.counts.setHorizontalHeaderLabels(['Country Code', 'Country name', 'Count'])
        self.counts.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        self.counts.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        layout.addWidget(self.counts, 1)
        self.source_versions = QLabel('No source versions loaded.')
        self.source_versions.setWordWrap(True)
        layout.addWidget(self.source_versions)
        export = QHBoxLayout()
        self.export_location = QLabel('No export yet.')
        self.export_location.setWordWrap(True)
        self.export_location.setTextInteractionFlags(Qt.TextInteractionFlag.TextSelectableByMouse)
        export.addWidget(self.export_location, 1)
        self.open_output_button = QPushButton('Open output folder')
        self.open_output_button.setEnabled(False)
        self.open_output_button.clicked.connect(self.open_output_folder)
        export.addWidget(self.open_output_button)
        layout.addLayout(export)
        self.log = QTextEdit()
        self.log.setReadOnly(True)
        self.log.setMaximumHeight(95)
        layout.addWidget(self.log)
        self.setCentralWidget(central)
        self.statusBar().showMessage('Ready for offline country analysis.')

    def _choose_input(self) -> None:
        path, _ = QFileDialog.getOpenFileName(
            self, 'Choose the address file', '', 'Address files (*.csv *.txt *.xlsx)')
        if path:
            self.input_path.setText(path)
            self.address_columns.clear()
            self.preview_input()

    def _choose_output(self) -> None:
        path = QFileDialog.getExistingDirectory(self, 'Choose the output folder')
        if path:
            self.output_dir.setText(path)

    def _choose_geo_db(self) -> None:
        path, _ = QFileDialog.getOpenFileName(
            self, 'Choose an installed GeoNames pack', '', 'SQLite packs (*.db *.sqlite)')
        if path:
            self.geo_db.setText(path)

    def _input_options(self) -> dict:
        return {'encoding': self.encoding.currentText(),
                'sheet_name': self.sheet_name.text().strip() or None,
                'address_columns': [item.text() for item in self.address_columns.selectedItems()]}

    def _update_column_selection(self) -> None:
        count = len(self.address_columns.selectedItems())
        self.column_selection.setText(
            f'{count} address columns selected. Click each address field to toggle selection; '
            'select all address lines, not page numbers or IDs.')

    def _notify(self, title: str, message: str) -> None:
        self.log.append(f'{title}: {message}')
        self.statusBar().showMessage(f'{title}: {message}')

    @property
    def busy(self) -> bool:
        # Include queued completion signals, so a second operation cannot replace
        # the callbacks before the first worker's finished slot is delivered.
        return self._worker_pending

    @property
    def job_state(self) -> str:
        return self._job_state

    @property
    def worker_thread(self):
        return self._thread

    @property
    def worker_thread_ident(self):
        return getattr(self._thread, 'thread_ident', None)

    def _start_worker(self, func, *, on_success=None, on_error=None) -> bool:
        if self.busy or self._closing:
            return False
        self._cancel.clear()
        thread = WorkerThread(func, cancel=self._cancel)
        thread.progressed.connect(self._on_progress)
        thread.succeeded.connect(self._worker_succeeded)
        thread.failed.connect(self._worker_failed)
        thread.finished.connect(partial(_thread_finished, thread))
        thread.finished.connect(self._worker_finished)
        self._thread = thread
        self._on_success = on_success
        self._on_error = on_error
        self._worker_pending = True
        self.inputs.setEnabled(False)
        for button in (self.run_button, self.resume_button, self.preview_button):
            button.setEnabled(False)
        self.cancel_button.setEnabled(True)
        self.progress.setRange(0, 0)
        _LIVE_THREADS.add(thread)
        thread.start()
        return True

    @Slot(object)
    def _worker_succeeded(self, payload) -> None:
        if not self._closing and self._on_success is not None:
            try:
                self._on_success(payload)
            except Exception as exc:
                self._operation_failed(f'{type(exc).__name__}: {exc}')

    @Slot(str)
    def _worker_failed(self, message: str) -> None:
        if not self._closing:
            callback = self._on_error or self._operation_failed
            callback(message)

    @Slot()
    def _worker_finished(self) -> None:
        self._on_success = None
        self._on_error = None
        self._worker_pending = False
        self.inputs.setEnabled(True)
        for button in (self.run_button, self.resume_button, self.preview_button):
            button.setEnabled(True)
        self.cancel_button.setEnabled(False)
        self.progress.setRange(0, 1)
        self.progress.setValue(int(self._job_state == 'done'))
        if self._closing:
            self.close()

    @Slot(int)
    def _on_progress(self, records: int) -> None:
        if not self._cancel.is_set():
            self.statusBar().showMessage(f'{records:,} records processed')

    def _operation_failed(self, message: str) -> None:
        self._last_error = message
        self._job_state = 'failed'
        self.summary.setText('Operation failed. See the error message below.')
        self._notify('Operation failed', message)
        self.jobFailed.emit(message)

    def preview_input(self) -> None:
        if self.busy or self._closing:
            return
        path = self.input_path.text().strip()
        if not path:
            self._notify('Missing input', 'Choose an address file first.')
            return
        self._load_columns(Path(path))

    def _load_columns(self, path: Path) -> None:
        options = self._input_options()
        cancel = self._cancel

        def work(_emit):
            return _preview_input(path, options, cancel.is_set)

        self._start_worker(work, on_success=self._apply_preview)

    def _apply_preview(self, payload: dict) -> None:
        if self._cancel.is_set() or (payload['preview'] or {}).get('cancelled'):
            self.preview_summary.setText('Preview cancelled. Counts are incomplete; preview again to finish.')
            self.statusBar().showMessage('Preview cancelled.')
            return
        self.address_columns.clear()
        for name in payload['header']:
            self.address_columns.addItem(name)
            self.address_columns.item(self.address_columns.count() - 1).setSelected(
                name in payload['selected'])
        if payload['sheet']:
            self.sheet_name.setText(payload['sheet'])
        preview = payload['preview']
        if preview is None:
            self.preview_summary.setText('Columns loaded. Select address columns, then preview again.')
            self.statusBar().showMessage('Select address columns before analyzing this input.')
        else:
            self.preview_summary.setText(
                f"Preview: {preview['total_records']:,} records; "
                f"{preview['empty_addresses']:,} empty addresses; "
                f"{preview['records_with_decoding_errors']:,} records with decoding errors.")
            self.statusBar().showMessage('Preview ready. No classification was run.')

    def start_job(self) -> None:
        if self.busy or self._closing:
            return
        source_text = self.input_path.text().strip()
        geo_text = self.geo_db.text().strip()
        if not source_text or not Path(source_text).is_file():
            self._notify('Missing input', 'Choose an existing address file first.')
            return
        if not geo_text or not Path(geo_text).is_file():
            self._notify('Missing GeoNames pack', 'Choose an installed GeoNames pack first.')
            return
        source = Path(source_text)
        options = self._input_options()
        if not options['address_columns']:
            self.summary.setText('Analysis not started. Select address columns first.')
            self._notify('Select address columns',
                         'Click Preview / load columns, then click each address field to select it. '
                         'Displayed column names are not automatically selected.')
            return
        chosen = self.output_dir.text().strip()
        output = Path(chosen) if chosen else source.parent / f'{source.stem}-country-analysis'
        try:
            occupied = output.exists() and (not output.is_dir() or next(output.iterdir(), None) is not None)
        except OSError as exc:
            self._operation_failed(f'Cannot inspect output folder: {exc}')
            return
        if occupied:
            self.summary.setText('Analysis not started. Choose a different empty output folder.')
            self._notify('Output folder is not empty',
                         'Existing files were not changed. Choose a new empty folder for a new analysis. '
                         'Use Resume only to continue the original run with its saved settings.')
            return
        geo_db = Path(geo_text)
        cancel = self._cancel

        def work(emit):
            from world_address_engine.jobs.country_analysis import analyze_job
            if not options['sheet_name'] and source.suffix.lower() == '.xlsx':
                _, options['sheet_name'] = _input_header(source, options['encoding'], None)
            result = analyze_job(source, output, geo_db, options,
                                 on_progress=emit, cancel_requested=cancel.is_set)
            if result.get('state', result.get('run_status')) == 'COMPLETE':
                return _read_country_export(output, result)
            return {'report': result, 'rows': []}

        self._begin_analysis(output, work)

    def resume_job(self) -> None:
        if self.busy or self._closing:
            return
        chosen = self.output_dir.text().strip()
        output = Path(chosen) if chosen else self._output_dir
        if output is None:
            self._notify('Nothing to resume', 'Choose the interrupted run output folder first.')
            return
        cancel = self._cancel

        def work(emit):
            from world_address_engine.jobs.country_analysis import resume_analysis
            result = resume_analysis(output, on_progress=emit, cancel_requested=cancel.is_set)
            if result.get('state', result.get('run_status')) == 'COMPLETE':
                return _read_country_export(output, result)
            return {'report': result, 'rows': []}

        self._begin_analysis(output, work)

    def _begin_analysis(self, output: Path, work) -> None:
        self._output_dir = output.resolve()
        self.output_dir.setText(str(self._output_dir))
        self.export_location.setText(f'Output folder: {self._output_dir}')
        self.open_output_button.setEnabled(False)
        self.counts.setRowCount(0)
        self.summary.setText('Analysis running...')
        self._job_state = 'running'
        self._last_error = ''
        self._last_report = None
        self.statusBar().showMessage('Analyzing countries offline...')
        self._start_worker(work, on_success=self._apply_analysis)

    def _apply_analysis(self, payload: dict) -> None:
        report = payload['report']
        if report.get('job_kind') != 'country-analysis':
            raise ValueError('Expected a country-analysis result.')
        state = report.get('state', report.get('run_status'))
        if state not in ('COMPLETE', 'PAUSED', 'CANCELLED'):
            raise ValueError(f'Unexpected analysis state: {state}')
        self._last_report = report
        self._job_state = {'COMPLETE': 'done', 'PAUSED': 'paused', 'CANCELLED': 'cancelled'}[state]
        counts = report.get('status_counts', {})
        breakdown = ' | '.join(
            f'{status.title()}: {counts[status]:,}' if status in counts
            else f'{status.title()}: pending'
            for status in ('DETERMINED', 'INFERRED', 'UNRESOLVED'))
        self.summary.setText(
            f"Total: {report.get('total_records', report.get('committed_records', 0)):,} | "
            f'{breakdown}')
        rows = payload['rows']
        self.counts.setRowCount(len(rows))
        for index, row in enumerate(rows):
            for column, value in enumerate(row):
                self.counts.setItem(index, column, QTableWidgetItem(str(value)))
        meta = report.get('run_metadata', {})
        self.source_versions.setText(
            f"Catalogue: {meta.get('catalogue_version', '')} | "
            f"GeoNames: {meta.get('geo_pack_version', meta.get('geo_pack_versions', ''))} | "
            f"Rules: {meta.get('rules_version', '')}")
        self.open_output_button.setEnabled(state == 'COMPLETE')
        message = ('Complete. Country reports are in the output folder.' if state == 'COMPLETE'
                   else f'{state.title()}. The checkpoint can be resumed from this output folder.')
        self._notify('Country analysis', message)
        self.jobFinished.emit()

    def cancel_job(self) -> None:
        self._cancel.set()
        self.cancel_button.setEnabled(False)
        if self.busy:
            self.statusBar().showMessage('Cancellation requested; waiting for the current operation...')

    def open_output_folder(self) -> None:
        if self._output_dir is not None:
            QDesktopServices.openUrl(QUrl.fromLocalFile(str(self._output_dir)))

    def closeEvent(self, event) -> None:
        if self.busy:
            self._closing = True
            self.cancel_job()
            self.statusBar().showMessage('Closing after the current operation stops...')
            event.ignore()
            return
        super().closeEvent(event)
