"""Offline, resumable country estimates with source-bound checkpoints."""
import csv
from contextlib import closing
import hashlib
import json
from pathlib import Path
import shutil
from time import monotonic

from world_address_engine.data.country_codes import load_country_codes
from world_address_engine.data.pack_registry import require_readable_pack
from world_address_engine.data.hk_regional import hk_pack_metadata, PROFILE as HK_PROFILE
from world_address_engine.io.country_reports import _private_directory, write_country_reports
from world_address_engine.io.delimited import _header_names, _record, iter_csv_records, iter_txt_records
from world_address_engine.jobs.checkpoint import (
    CHUNK_PATTERN, HEADER_KEY, atomic_json, commit_batch, committed_shards, file_sha256,
)
from world_address_engine.matching.country import (
    COUNTRY_RULES_VERSION, CountryClassifier, CountryInference,
)
from world_address_engine.models import AddressRecord
from world_address_engine.sqlite_util import open_db

STATE_VERSION = '2026-10-26.country-state.3'
JOB_KIND = 'country-analysis'
FINALIZATION_DIR = '.country-finalize'


class _Cancelled(Exception):
    pass


class _StageProgress:
    """Bound notifications without retaining records or delaying cancellation."""

    def __init__(self, callback):
        self.callback = callback
        self.stage = None
        self.processed = 0
        self.started_at = 0
        self.reported_work = False
        self.last_time = 0.0

    def emit(self, stage, processed=0, committed=0, *, total=None,
             status_counts=None, force=False):
        changed = stage != self.stage
        if changed:
            self.started_at = processed
            self.reported_work = False
        self.stage, self.processed = stage, processed
        if self.callback is None:
            return
        now = monotonic()
        first = not self.reported_work and processed > self.started_at
        if not (changed or first or force or now - self.last_time >= 0.25):
            return
        self.last_time = now
        self.reported_work |= first
        event = {'stage': stage, 'processed': processed, 'committed': committed}
        if total is not None:
            event['total'] = total
        if status_counts is not None:
            event['status_counts'] = dict(status_counts)
        self.callback(event)


def _json_hash(value) -> str:
    return hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=True,
                                     separators=(',', ':')).encode('ascii')).hexdigest()


def _options(options: dict | None) -> dict:
    result = {'encoding': 'utf-8-sig', 'address_columns': [], 'record_id_column': None,
              'sheet_name': None, 'txt_mode': 'line', 'delimiter': None}
    if options:
        if set(options) - result.keys():
            raise ValueError('unsupported country analysis input options')
        result.update(options)
    result['address_columns'] = list(result['address_columns'] or ())
    return result


def _records(path: Path, options: dict):
    kind = path.suffix.lower()
    columns = tuple(options['address_columns'])
    if kind == '.csv':
        return iter_csv_records(path, options['encoding'], columns, options['record_id_column'])
    if kind == '.txt':
        return iter_txt_records(path, options['encoding'], options['txt_mode'], options['delimiter'])
    if kind == '.xlsx':
        from world_address_engine.io.xlsx import iter_xlsx_records
        if not options['sheet_name']:
            raise ValueError('Excel input requires sheet_name')
        return iter_xlsx_records(path, options['sheet_name'], columns, options['record_id_column'])
    raise ValueError('expected CSV, XLSX, or TXT input')


def _original_columns(path: Path, options: dict) -> list[str]:
    """Read just the header so a zero-record input keeps its original schema."""
    if path.suffix.lower() == '.csv':
        with path.open(encoding=options['encoding'], errors='surrogateescape', newline='') as stream:
            return _header_names(next(csv.reader(stream), []))
    if path.suffix.lower() == '.xlsx':
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            if options['sheet_name'] not in book.sheetnames:
                raise ValueError(f"worksheet not found: {options['sheet_name']}")
            headers = next(book[options['sheet_name']].iter_rows(values_only=True), ())
            return _header_names(['' if value is None else str(value) for value in headers])
        finally:
            book.close()
    return []


def _preflight_input(path: Path, options: dict) -> None:
    """Validate mappings even without data rows, without consuming the job stream."""
    if path.suffix.lower() in ('.csv', '.xlsx'):
        if path.suffix.lower() == '.xlsx' and not options['sheet_name']:
            raise ValueError('Excel input requires sheet_name')
        headers = _original_columns(path, options)
        if not headers:
            raise ValueError('input has no header')
        # The readers share this mapping validation, but invoke it only per row.
        _record(1, path, dict.fromkeys(headers, ''), tuple(options['address_columns']),
                options['record_id_column'])
    else:
        # TXT option checks are lazy; advance once and always close the probe.
        with closing(_records(path, options)) as records:
            next(records, None)


def _pack_metadata(geo_db: Path, geo_digest: str) -> dict:
    manifest_path = geo_db.with_name('manifest.json')
    manifest = json.loads(manifest_path.read_text(encoding='utf8')) if manifest_path.is_file() else {}
    if 'db_sha256' in manifest and manifest['db_sha256'] != geo_digest:
        raise ValueError('geo pack checksum does not match its manifest')
    if manifest.get('profile') == HK_PROFILE:
        raise ValueError('Choose a global GeoNames pack here; Hong Kong regional packs have a separate option.')
    with open_db(geo_db) as connection:
        tables = {row[0] for row in connection.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        if 'hk_pack_metadata' in tables:
            raise ValueError('Choose a global GeoNames pack here; Hong Kong regional packs have a separate option.')
        required_columns = {
            'geo_places': set('place_id name country_code admin1_code admin2_code feature_class feature_code population'.split()),
            'geo_aliases': {'place_id', 'alias', 'search_key', 'language'},
        }
        for name, columns in required_columns.items():
            actual = {row[1] for row in connection.execute(f'PRAGMA table_info({name})')}
            if name not in tables or not columns <= actual:
                raise ValueError('The selected database is not a readable global GeoNames pack')
        table = connection.execute(
            "SELECT 1 FROM sqlite_master WHERE type='table' AND name='pack_metadata'").fetchone()
        row = connection.execute('SELECT profile FROM pack_metadata LIMIT 1').fetchone() if table else None
    profile = str(row[0]) if row else str(manifest.get('profile') or 'UNKNOWN')
    if manifest.get('profile') and manifest['profile'] != profile:
        raise ValueError('geo pack profile does not match its manifest')
    generation = require_readable_pack(geo_db)
    return {'profile': profile, 'pack_format_version': generation,
            'geo_pack_version': str(manifest.get('pack_version') or 'sha256:' + geo_digest),
            'geo_manifest_sha256': file_sha256(manifest_path) if manifest_path.is_file() else None}


def _catalogue_metadata(catalogue) -> dict:
    return {'catalogue_version': catalogue.version,
            'catalogue_sha256': catalogue.source_sha256,
            'catalogue_fingerprint': catalogue.fingerprint}


def _load_zones():
    from world_address_engine.data.zone_mappings import load_zone_mappings
    return load_zone_mappings()


def _zone_metadata(zones=None) -> dict:
    zones = _load_zones() if zones is None else zones
    return {'version': zones.version, 'source_sha256': zones.source_sha256,
            'sha256': zones.fingerprint}


def _assert_sources_unchanged(cfg: dict, *, on_event=None) -> None:
    if cfg.get('job_kind') != JOB_KIND or cfg.get('schema_version') != STATE_VERSION:
        if cfg.get('schema_version') == '2026-10-24.country-state.2':
            raise ValueError('Use v0.2.0 to resume this older country checkpoint; start a new analysis for zone output.')
        raise ValueError('unsupported country checkpoint version or kind')
    if cfg.get('rules_version') != COUNTRY_RULES_VERSION:
        raise ValueError('country rules version changed; start a new analysis')
    if cfg.get('zone_mapping') != _zone_metadata():
        raise ValueError('zone mapping changed since country checkpoint')
    if _json_hash(cfg['input_options']) != cfg['input_options_sha256']:
        raise ValueError('input options profile changed since checkpoint')
    for label, path_key, hash_key in (('input', 'input_path', 'input_sha256'),
                                       ('geo pack', 'geo_db', 'geo_sha256')):
        if on_event is not None:
            on_event({'stage': 'hashing_input' if label == 'input' else 'hashing_pack',
                      'processed': 0, 'committed': 0})
        if file_sha256(Path(cfg[path_key])) != cfg[hash_key]:
            raise ValueError(f'{label} changed since country checkpoint')
    if on_event is not None:
        on_event({'stage': 'loading', 'processed': 0, 'committed': 0})
    for key, value in _pack_metadata(Path(cfg['geo_db']), cfg['geo_sha256']).items():
        if cfg.get(key) != value:
            raise ValueError('geo pack version or profile changed since checkpoint')
    regional = cfg.get('hk_regional')
    if regional is not None:
        if on_event is not None:
            on_event({'stage': 'hashing_hk_pack', 'processed': 0, 'committed': 0})
        try:
            current = hk_pack_metadata(Path(regional['path']))
        except (KeyError, OSError, ValueError) as exc:
            raise ValueError('Hong Kong regional pack changed or is unavailable since checkpoint') from exc
        if current != regional:
            raise ValueError('Hong Kong regional pack changed since checkpoint')
    catalogue = load_country_codes(Path(cfg['catalogue_path']) if cfg['catalogue_path'] else None)
    if any(cfg.get(key) != value for key, value in _catalogue_metadata(catalogue).items()):
        raise ValueError('country catalogue changed since checkpoint')


def _saved_pairs(paths, rules_version):
    """Validate the bytes being consumed, not just an earlier checksum pass."""
    previous = 0
    previous_digest = None
    for path in paths:
        match = CHUNK_PATTERN.fullmatch(path.name)
        if not match:
            raise ValueError('invalid country checkpoint filename')
        start, end, expected_digest = int(match[1]), int(match[2]), match[3]
        digest = hashlib.sha256()
        count = 0
        with path.open('rb') as stream:
            for index, raw in enumerate(stream):
                digest.update(raw)
                try:
                    payload = json.loads(raw)
                    if index == 0:
                        if payload != {HEADER_KEY: {'start': start, 'end': end,
                                                   'previous': previous_digest,
                                                   'rules_version': rules_version}}:
                            raise ValueError('country checkpoint header mismatch')
                        continue
                    record = AddressRecord(**payload['record'])
                    result_data = dict(payload['result'])
                    for key in ('evidence', 'candidates'):
                        result_data[key] = tuple(result_data.get(key, ()))
                    inference = CountryInference(**result_data)
                except (KeyError, TypeError, ValueError) as exc:
                    raise ValueError('invalid country checkpoint row or header') from exc
                if record.record_no != previous + 1 or inference.record_no != record.record_no:
                    raise ValueError('country checkpoint record ordering mismatch')
                previous += 1
                count += 1
                yield record, inference
        if count != end - start + 1 or previous != end or start != previous - count + 1:
            raise ValueError('country checkpoint row count or boundary mismatch')
        if digest.hexdigest() != expected_digest:
            raise ValueError('country checkpoint checksum mismatch')
        previous_digest = expected_digest


def _verified_pairs(paths, source, classifier, *, complete=False, cancel_requested=None,
                    on_progress=None):
    processed = 0
    with closing(_saved_pairs(paths, COUNTRY_RULES_VERSION)) as saved:
        for saved_record, saved_result in saved:
            if cancel_requested is not None and cancel_requested():
                raise _Cancelled
            original = next(source, None)
            if original is None or original != saved_record:
                raise ValueError('country checkpoint verification failed: original source record mismatch')
            fresh = classifier.classify(original)
            if fresh != saved_result:
                raise ValueError('country checkpoint verification failed: fresh inference mismatch')
            processed += 1
            if on_progress is not None:
                on_progress(processed)
            yield original, fresh
    if complete and next(source, None) is not None:
        raise ValueError('country checkpoint verification failed: omitted source records')


def _remove_checkpoint(state: Path) -> None:
    if state.exists():
        shutil.rmtree(state)
    if state.exists():
        raise OSError('country checkpoint cleanup failed')


def _limits(batch_size: int, max_batches: int | None) -> None:
    if isinstance(batch_size, bool) or not isinstance(batch_size, int) or not 1 <= batch_size <= 100_000:
        raise ValueError('batch_size must be 1 to 100000')
    if max_batches is not None and (
            isinstance(max_batches, bool) or not isinstance(max_batches, int) or max_batches < 1):
        raise ValueError('max_batches must be positive')


def _process(output: Path, cfg: dict, *, max_batches=None, on_progress=None, on_event=None,
             cancel_requested=None, keep_checkpoint=False) -> dict:
    progress = _StageProgress(on_event)
    progress.emit('loading')
    state = output / '.country-state'
    paths, last_committed = committed_shards(state, expected_rules_version=COUNTRY_RULES_VERSION)
    catalogue = load_country_codes(Path(cfg['catalogue_path']) if cfg['catalogue_path'] else None)
    columns = _original_columns(Path(cfg['input_path']), cfg['input_options'])
    zones = _load_zones()
    if cfg.get('zone_mapping') != _zone_metadata(zones):
        raise ValueError('zone mapping changed since country checkpoint')
    source = iter(_records(Path(cfg['input_path']), cfg['input_options']))
    sparse_columns = Path(cfg['input_path']).suffix.lower() == '.txt'
    batch = []
    batches = 0

    def check_cancel():
        if cancel_requested is not None and cancel_requested():
            raise _Cancelled

    def commit():
        nonlocal last_committed, batches
        commit_batch(state, batch, COUNTRY_RULES_VERSION)
        last_committed = batch[-1][0].record_no
        batches += 1
        batch.clear()
        if on_progress is not None:
            on_progress(last_committed)

    try:
        check_cancel()
        regional_path = Path(cfg['hk_regional']['path']) if cfg.get('hk_regional') else None
        with CountryClassifier(Path(cfg['geo_db']), catalogue, hk_regional_db=regional_path) as classifier:
            if classifier.rules_version != cfg['rules_version']:
                raise ValueError('country classifier rules version changed')
            # Restore only after comparing each committed record and inference.
            progress.emit('restoring', committed=last_committed, total=last_committed)
            with closing(_verified_pairs(paths, source, classifier,
                                         cancel_requested=cancel_requested,
                                         on_progress=lambda count: progress.emit(
                                             'restoring', count, last_committed,
                                             total=last_committed))) as restored:
                for record, _ in restored:
                    if sparse_columns and len(record.fields) > len(columns):
                        columns = list(record.fields)
            progress.emit('restoring', last_committed, last_committed,
                          total=last_committed, force=True)
            progress.emit('classifying', last_committed, last_committed)
            for record in source:
                check_cancel()
                if record.record_no != last_committed + len(batch) + 1:
                    raise ValueError('country source record ordering mismatch')
                if sparse_columns and len(record.fields) > len(columns):
                    columns = list(record.fields)
                batch.append((record, classifier.classify(record)))
                progress.emit('classifying', record.record_no, last_committed)
                if len(batch) >= cfg['batch_size']:
                    commit()
                    check_cancel()
                    if max_batches is not None and batches >= max_batches:
                        progress.emit('classifying', last_committed, last_committed, force=True)
                        progress.emit('paused', last_committed, last_committed)
                        return {'run_status': 'PAUSED', 'job_kind': JOB_KIND,
                                'committed_records': last_committed}
            check_cancel()
            if batch:
                commit()
            progress.emit('classifying', last_committed, last_committed,
                          total=last_committed, force=True)
            progress.emit('verifying_sources', last_committed, last_committed,
                          total=last_committed)
            check_cancel()
            _assert_sources_unchanged(cfg)
            paths, last_committed = committed_shards(state, expected_rules_version=COUNTRY_RULES_VERSION)
            metadata = {key: cfg[key] for key in (
                'rules_version', 'catalogue_version', 'catalogue_sha256', 'catalogue_fingerprint',
                'input_sha256', 'geo_sha256', 'geo_manifest_sha256', 'geo_pack_version',
                'pack_format_version', 'profile', 'input_options_sha256', 'batch_size')}
            metadata['address_columns'] = cfg['input_options']['address_columns']
            metadata['hk_regional'] = cfg.get('hk_regional')
            metadata['zone_mapping'] = dict(cfg['zone_mapping'])

            def before_manifest():
                progress.emit('exporting', last_committed, last_committed,
                              total=last_committed, force=True)
                progress.emit('finalizing', last_committed, last_committed,
                              total=last_committed)
                _assert_sources_unchanged(cfg)
                check_cancel()
                if not keep_checkpoint:
                    # Keep only pinned configuration across the gap between
                    # deleting plaintext shards and publishing the manifest.
                    recovery = output / FINALIZATION_DIR
                    recovery.mkdir(mode=0o700, exist_ok=True)
                    _private_directory(recovery)
                    atomic_json(recovery / 'config.json', cfg)
                    _remove_checkpoint(state)

            progress.emit('exporting', committed=last_committed, total=last_committed)
            with closing(iter(_records(Path(cfg['input_path']), cfg['input_options']))) as originals, \
                    closing(_verified_pairs(paths, originals, classifier, complete=True,
                                            cancel_requested=cancel_requested,
                                            on_progress=lambda count: progress.emit(
                                                'exporting', count, last_committed,
                                                total=last_committed))) as verified:
                result = write_country_reports(
                    verified, output, catalogue, metadata,
                    original_columns=columns, allow_sparse_columns=sparse_columns,
                    before_manifest=before_manifest, zone_catalogue=zones)
            recovery = output / FINALIZATION_DIR
            if recovery.exists():
                shutil.rmtree(recovery)
            progress.emit('complete', last_committed, last_committed, total=last_committed,
                          status_counts=result['status_counts'])
            return result
    except _Cancelled:
        progress.emit('cancelled', progress.processed, last_committed)
        return {'run_status': 'CANCELLED', 'job_kind': JOB_KIND,
                'committed_records': last_committed}
    finally:
        source.close()


def analyze_job(input_path, output_dir, geo_db, input_options: dict | None = None, *,
                catalogue_path=None, hk_regional_db=None, batch_size=1000, max_batches=None,
                on_progress=None, on_event=None, cancel_requested=None, keep_checkpoint=False,
                expected_reference_identity=None) -> dict:
    """Start an estimate with integer commit progress and optional count-only events.

    A new job requires an empty output directory. PAUSED and CANCELLED results
    leave source-bound sensitive state for resume_analysis; COMPLETE removes it
    unless keep_checkpoint is explicitly enabled. Event processed counts belong
    to the current stage; total is omitted until known without a counting pass.
    """
    progress = _StageProgress(on_event)
    progress.emit('validating')
    input_path, output, geo_db = map(Path, (input_path, output_dir, geo_db))
    if output.exists() and any(output.iterdir()):
        raise FileExistsError('existing output directory; resume analysis or choose an empty directory')
    _limits(batch_size, max_batches)
    options = _options(input_options)
    for path in (input_path, geo_db):
        if not path.is_file():
            raise FileNotFoundError(path)
    input_path, geo_db = input_path.resolve(), geo_db.resolve()
    _preflight_input(input_path, options)
    progress.emit('hashing_input')
    input_digest = file_sha256(input_path)
    progress.emit('hashing_pack')
    geo_digest = file_sha256(geo_db)
    regional = None
    if hk_regional_db is not None:
        progress.emit('hashing_hk_pack')
        regional = hk_pack_metadata(Path(hk_regional_db))
    progress.emit('loading')
    catalogue_path = Path(catalogue_path).resolve() if catalogue_path is not None else None
    catalogue = load_country_codes(catalogue_path)
    zones = _load_zones()
    cfg = {'job_kind': JOB_KIND, 'schema_version': STATE_VERSION,
           'hk_regional': regional, 'zone_mapping': _zone_metadata(zones),
           'rules_version': COUNTRY_RULES_VERSION, 'batch_size': batch_size,
           'input_path': str(input_path), 'geo_db': str(geo_db),
           'catalogue_path': str(catalogue_path) if catalogue_path is not None else None,
           'input_sha256': input_digest, 'geo_sha256': geo_digest,
           'input_options': options, 'input_options_sha256': _json_hash(options),
           **_pack_metadata(geo_db, geo_digest), **_catalogue_metadata(catalogue)}
    if expected_reference_identity is not None:
        actual_identity = {key: cfg[key] for key in
                           ('geo_sha256', 'geo_manifest_sha256', 'hk_regional')}
        if expected_reference_identity != actual_identity:
            raise ValueError('reference packs changed after verification')
    output.mkdir(parents=True, exist_ok=True)
    if any(output.iterdir()):
        raise FileExistsError('existing country analysis output')
    state = output / '.country-state'
    state.mkdir(mode=0o700)
    _private_directory(state)
    atomic_json(state / 'config.json', cfg)
    return _process(output, cfg, max_batches=max_batches, on_progress=on_progress,
                    on_event=on_event,
                    cancel_requested=cancel_requested, keep_checkpoint=keep_checkpoint)


def resume_analysis(output_dir, *, max_batches=None, on_progress=None, on_event=None,
                    cancel_requested=None, keep_checkpoint=False) -> dict:
    """Resume with mandatory source, catalogue, pack, profile and inference checks."""
    progress = _StageProgress(on_event)
    progress.emit('validating')
    output = Path(output_dir)
    if (output / 'Run_Manifest.json').exists():
        raise FileExistsError('completed country analysis is immutable')
    state = output / '.country-state'
    config_path = state / 'config.json'
    restore_config = not config_path.is_file()
    if restore_config:
        config_path = output / FINALIZATION_DIR / 'config.json'
    cfg = json.loads(config_path.read_text(encoding='utf8'))
    _limits(cfg['batch_size'], max_batches)
    _assert_sources_unchanged(cfg, on_event=on_event)
    if restore_config:
        # A failed finalization has no authoritative row checkpoint left.
        # Rebuild and verify every report from the unchanged original source.
        _remove_checkpoint(state)
        state.mkdir(mode=0o700)
    _private_directory(state)
    if restore_config:
        atomic_json(state / 'config.json', cfg)
    return _process(output, cfg, max_batches=max_batches, on_progress=on_progress,
                    on_event=on_event,
                    cancel_requested=cancel_requested, keep_checkpoint=keep_checkpoint)
