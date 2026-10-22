"""Memory-bounded XLSX address reader with required-field formula checks."""
from itertools import zip_longest
from pathlib import Path
from collections.abc import Iterator
from openpyxl import load_workbook
from world_address_engine.io.delimited import _header_names, _record
from world_address_engine.models import AddressRecord


def iter_xlsx_records(path: Path, sheet_name: str, address_columns: tuple[str,...],
                      record_id_column: str | None = None) -> Iterator[AddressRecord]:
    path=Path(path)
    values_book=load_workbook(path,read_only=True,data_only=True)
    formulas_book=load_workbook(path,read_only=True,data_only=False)
    try:
        if sheet_name not in values_book.sheetnames:
            raise ValueError(f'worksheet not found: {sheet_name}')
        value_rows=values_book[sheet_name].iter_rows(values_only=True)
        formula_rows=formulas_book[sheet_name].iter_rows(values_only=True)
        headers=next(value_rows,None)
        next(formula_rows,None)
        if not headers:
            raise ValueError('Excel sheet has no header')
        keys=_header_names([str(c) if c is not None else '' for c in headers])
        columns=address_columns or tuple(f'NAME_ADDRESS_LINE{i}' for i in range(1,6)
                                          if f'NAME_ADDRESS_LINE{i}' in keys)
        if not columns:
            raise ValueError('No Excel address columns detected or specified')
        seq=0
        for physical_row,(original,cells_formula) in enumerate(zip_longest(value_rows,formula_rows,fillvalue=()),2):
            if not original or all(x is None for x in original) and all(x is None for x in cells_formula):
                continue
            seq+=1
            data={key: '' if j>=len(original) or original[j] is None else str(original[j])
                  for j,key in enumerate(keys)}
            record=_record(seq,path,data,tuple(columns),record_id_column)
            formula_missing=any(j<len(cells_formula) and isinstance(cells_formula[j],str)
                                and cells_formula[j].startswith('=')
                                and (j>=len(original) or original[j] is None)
                                and key in columns
                                for j,key in enumerate(keys))
            yield AddressRecord(record.record_no,record.source_file,sheet_name,
                                record.record_id,record.fields,record.address_text,
                                record.country_hint,
                                'FORMULA_WITHOUT_CACHED_VALUE' if formula_missing else record.input_error)
    finally:
        values_book.close()
        formulas_book.close()
