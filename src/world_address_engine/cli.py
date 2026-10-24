"""Offline address input preview and conservative TEST-ONLY batch tooling.

The preview command needs no data packs. Classification fails closed until an
independently audited Hongkong Post reference master becomes available.
"""
import argparse
from collections import Counter
import json
from pathlib import Path
import sys
from typing import Callable
from world_address_engine.io.delimited import iter_csv_records, iter_txt_records


def preview(path: Path, encoding: str = 'utf-8-sig', address_columns: tuple[str,...] = (),
            sheet_name: str | None = None, txt_mode: str = 'line',
            delimiter: str | None = None, *,
            cancel_requested: Callable[[], bool] | None = None) -> dict:
    kind=path.suffix.lower()
    if kind=='.csv':
        rows=iter_csv_records(path,encoding,address_columns)
    elif kind=='.txt':
        rows=iter_txt_records(path,encoding,txt_mode,delimiter)
    elif kind=='.xlsx':
        from openpyxl import load_workbook
        from world_address_engine.io.xlsx import iter_xlsx_records
        if not sheet_name:
            wb=load_workbook(path,read_only=True)
            try:sheet_name=wb.sheetnames[0]
            finally:wb.close()
        rows=iter_xlsx_records(path,sheet_name,address_columns)
    else:
        raise ValueError('Unsupported input format: use .csv, .xlsx or .txt')
    total=0;issues=0; extras=0; empty=0; lines=Counter()
    headers=[]; used=[]
    cancelled = False
    try:
        for record in rows:
            if cancel_requested is not None and cancel_requested():
                cancelled = True
                break
            total+=1
            if total==1:
                headers=list(record.fields)
                used=list(address_columns) if address_columns else [x for x in headers if x.startswith('NAME_ADDRESS_LINE')]
                if not used and kind=='.txt': used=[x for x in headers if x.startswith('Text_')]
            issues+=bool(record.input_error)
            extras+=any(v for k,v in record.fields.items() if k.startswith('Unnamed_'))
            empty+=not bool(record.address_text.strip())
            lines[len([v for v in record.address_text.splitlines() if v.strip()])]+=1
    finally:
        rows.close()
    result = {
        'input_file':path.name, 'format':kind.lstrip('.'), 'total_records':total,
        'original_columns':len(headers), 'address_columns':used,
        'records_with_extra_unnamed_data':extras,
        'records_with_decoding_errors':issues, 'empty_addresses':empty,
        'nonempty_address_line_distribution':dict(sorted(lines.items())),
        'classifications_run':False,
    }
    if cancelled:
        result['cancelled'] = True
    return result


def main(argv: list[str] | None=None) -> int:
    parser=argparse.ArgumentParser(description='Fully offline global postal address tooling')
    sub=parser.add_subparsers(dest='command',required=True)
    scan=sub.add_parser('preview',help='Show input shape and counts without printing addresses')
    scan.add_argument('--input',type=Path,required=True)
    scan.add_argument('--encoding',default='utf-8-sig')
    scan.add_argument('--address-column',action='append',default=[])
    scan.add_argument('--sheet')
    scan.add_argument('--txt-mode',choices=['line','delimited'],default='line')
    scan.add_argument('--delimiter')
    classify = sub.add_parser('classify', help='Run offline classification (audited master required; --test-only for synthetic fixtures)')
    classify.add_argument('--input', type=Path, required=True)
    classify.add_argument('--output', type=Path, required=True)
    classify.add_argument('--geo-db', type=Path, required=True)
    classify.add_argument('--hkp-db', type=Path, required=True)
    classify.add_argument('--encoding', default='utf-8-sig')
    classify.add_argument('--address-column', action='append', default=[])
    classify.add_argument('--record-id-column')
    classify.add_argument('--sheet')
    classify.add_argument('--txt-mode', choices=['line', 'delimited'], default='line')
    classify.add_argument('--delimiter')
    classify.add_argument('--batch-size', type=int, default=10000)
    classify.add_argument('--max-batches', type=int)
    classify.add_argument('--approval-manifest', type=Path,
                          help='Signed HKP approval manifest (required outside --test-only)')
    classify.add_argument('--approval-key', type=Path,
                          help='Key for the approval manifest; keep it outside the repository')
    classify.add_argument('--test-only', action='store_true', help='Explicitly mark the output TEST ONLY; no production postal use')
    classify.add_argument('--verify-shards', action='store_true',
                          help='Deprecated compatibility flag; verification is always enabled')
    classify.add_argument('--keep-checkpoint', action='store_true',
                        help='Keep the plaintext customer checkpoint shards after the run')
    analyze = sub.add_parser('analyze', help='Estimate country codes fully offline without an HKP database')
    analyze.add_argument('--input', type=Path, required=True)
    analyze.add_argument('--output', type=Path, required=True)
    analyze.add_argument('--geo-db', type=Path, required=True)
    analyze.add_argument('--catalogue', type=Path)
    analyze.add_argument('--encoding', default='utf-8-sig')
    analyze.add_argument('--address-column', action='append', default=[])
    analyze.add_argument('--record-id-column')
    analyze.add_argument('--sheet')
    analyze.add_argument('--txt-mode', choices=['line', 'delimited'], default='line')
    analyze.add_argument('--delimiter')
    analyze.add_argument('--batch-size', type=int, default=1000)
    analyze.add_argument('--max-batches', type=int)
    analyze.add_argument('--keep-checkpoint', action='store_true',
                         help='Keep sensitive plaintext country checkpoint shards after completion')
    resume = sub.add_parser('resume', help='Resume a paused run with source and data-pack integrity checks')
    resume.add_argument('--job-dir', type=Path, required=True)
    resume.add_argument('--max-batches', type=int)
    resume.add_argument('--approval-manifest', type=Path)
    resume.add_argument('--approval-key', type=Path)
    resume.add_argument('--verify-shards', action='store_true',
                        help='Deprecated compatibility flag; verification is always enabled')
    resume.add_argument('--keep-checkpoint', action='store_true',
                        help='Keep the plaintext customer checkpoint shards after the run')
    args=parser.parse_args(argv)
    try:
        if args.command == 'preview':
            report=preview(args.input,args.encoding,tuple(args.address_column),args.sheet,args.txt_mode,args.delimiter)
        elif args.command == 'classify':
            from world_address_engine.jobs.runner import run_job
            options = {'encoding':args.encoding,'address_columns':args.address_column,
                       'record_id_column':args.record_id_column,'sheet_name':args.sheet,
                       'txt_mode':args.txt_mode,'delimiter':args.delimiter}
            report=run_job(args.input,args.output,args.geo_db,args.hkp_db,options,
                           batch_size=args.batch_size,max_batches=args.max_batches,
                           allow_unverified_reference_data=args.test_only,
                           approval_manifest=args.approval_manifest,
                           approval_key=args.approval_key,
                           keep_checkpoint=args.keep_checkpoint)
        elif args.command == 'analyze':
            from world_address_engine.jobs.country_analysis import analyze_job
            options = {'encoding': args.encoding, 'address_columns': args.address_column,
                       'record_id_column': args.record_id_column, 'sheet_name': args.sheet,
                       'txt_mode': args.txt_mode, 'delimiter': args.delimiter}
            report = analyze_job(args.input, args.output, args.geo_db, options,
                                 catalogue_path=args.catalogue, batch_size=args.batch_size,
                                 max_batches=args.max_batches, keep_checkpoint=args.keep_checkpoint)
        elif any((args.job_dir / name).is_dir()
                 for name in ('.country-state', '.country-finalize')) or (
                (args.job_dir / 'Run_Manifest.json').is_file()
                and json.loads((args.job_dir / 'Run_Manifest.json').read_text(encoding='utf8'))
                    .get('job_kind') == 'country-analysis'):
            from world_address_engine.jobs.country_analysis import resume_analysis
            report = resume_analysis(args.job_dir, max_batches=args.max_batches,
                                     keep_checkpoint=args.keep_checkpoint)
        else:
            from world_address_engine.jobs.runner import resume_job
            report=resume_job(args.job_dir,max_batches=args.max_batches,
                              keep_checkpoint=args.keep_checkpoint or None)
    except (ValueError,OSError,LookupError) as exc:
        print(f'Input preview failed: {exc}',file=sys.stderr)
        return 1
    print(json.dumps(report,ensure_ascii=False,indent=2))
    return 0


if __name__=='__main__':
    raise SystemExit(main())
